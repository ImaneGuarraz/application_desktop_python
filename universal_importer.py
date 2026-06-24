# universal_importer.py

"""
Ce fichier contient un importeur universel capable de lire plusieurs formats : json, csv, txt, docx, xlsx.
Son rôle est d'extraire automatiquement des paires (name, value) même si la structure du fichier est complexe ou inconnue.
Il est utilisé par service.py pour insérer des données dans sqlite """

import json
import csv
import docx          # pour lire les fichiers .docx
import openpyxl      # pour lire les fichiers excel


# détecte automatiquement le type de fichier selon son extension
def detect_file_type(path):
    # récupère l'extension du fichier
    ext = path.lower().split(".")[-1]

    # json
    if ext == "json":
        return "json"

    # csv ou txt
    if ext in ("csv", "txt"):
        return "csv"

    # word
    if ext == "docx":
        return "docx"

    # excel
    if ext in ("xlsx", "xls"):
        return "excel"

    # format inconnu
    return "unknown"


# aplatit un json complexe en liste de paires (clé, valeur)
def flatten_json(obj, parent_key="", sep="."):
    # liste des résultats
    items = []

    # si l'objet est un dictionnaire
    if isinstance(obj, dict):
        # parcourt chaque clé
        for k, v in obj.items():
            # construit la clé complète
            new_key = f"{parent_key}{sep}{k}" if parent_key else k
            # aplatit récursivement
            items.extend(flatten_json(v, new_key, sep=sep))

    # si l'objet est une liste
    elif isinstance(obj, list):
        # parcourt chaque élément
        for i, v in enumerate(obj):
            new_key = f"{parent_key}{sep}{i}" if parent_key else str(i)
            items.extend(flatten_json(v, new_key, sep=sep))

    # si l'objet est une valeur simple
    else:
        items.append((parent_key, obj))

    return items


# choisit automatiquement un champ texte + un champ numérique
def auto_pick_name_value(flat_items):
    # stocke la première valeur texte trouvée
    text_key = None

    # stocke la première valeur numérique trouvée
    num_key = None

    # parcourt toutes les paires (clé, valeur)
    for key, value in flat_items:

        # si la valeur est un texte et qu'on n'en a pas encore
        if isinstance(value, str) and not text_key:
            text_key = (key, value)

        # si la valeur est un nombre et qu'on n'en a pas encore
        if isinstance(value, (int, float)) and not num_key:
            num_key = (key, value)

        # si on a trouvé les deux, on peut retourner
        if text_key and num_key:
            return text_key[1], float(num_key[1])

    # si rien trouvé, retourne None
    return None, None


# extraction universelle selon le type de fichier
def extract_data_any_format(path):
    # détecte le type du fichier
    file_type = detect_file_type(path)

    # liste des résultats (name, value)
    results = []

    ####### JSON
    if file_type == "json":
        # ouvre le fichier json
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)

        # si le json contient "records" ou "data", on les utilise
        if isinstance(data, dict):
            data = data.get("records", data.get("data", data))

        # si c'est encore un dict, on le met dans une liste
        if isinstance(data, dict):
            data = [data]

        # parcourt chaque objet json
        for obj in data:
            flat = flatten_json(obj)               # aplatit le json
            name, value = auto_pick_name_value(flat)  # choisit name + value
            if name and value is not None:
                results.append((name, value))

    ####### CSV / TXT
    elif file_type == "csv":
        # ouvre le fichier csv ou txt
        with open(path, "r", encoding="utf-8") as f:
            reader = csv.reader(f)

            # parcourt chaque ligne
            for row in reader:
                if len(row) < 2:
                    continue

                name = row[0].strip()

                # tente de convertir la deuxième colonne en nombre
                try:
                    value = float(row[1])
                    results.append((name, value))
                except:
                    continue

    ####### DOCX
    elif file_type == "docx":
        # ouvre le fichier word
        doc = docx.Document(path)

        # parcourt chaque paragraphe
        for para in doc.paragraphs:
            parts = para.text.split()

            # on attend au moins deux mots : name + value
            if len(parts) >= 2:
                name = parts[0]

                try:
                    value = float(parts[1])
                    results.append((name, value))
                except:
                    continue

    ####### EXCEL
    elif file_type == "excel":
        # ouvre le fichier excel
        wb = openpyxl.load_workbook(path)
        ws = wb.active

        # parcourt chaque ligne
        for row in ws.iter_rows(values_only=True):
            if len(row) < 2:
                continue

            name = str(row[0])

            try:
                value = float(row[1])
                results.append((name, value))
            except:
                continue

    # retourne la liste finale
    return results
